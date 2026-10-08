"""Regenerates the workbook fixtures of this directory.

    PYTHONPATH=<dir holding xlwt> python3 ops/fixtures/workbooks/make_workbooks.py

Needs `openpyxl` (writes `.xlsx`) and `xlwt` (writes `.xls`, BIFF8). Neither is
a dependency of the repo: install them wherever you like, or unpack the `xlwt`
wheel into a scratch directory and put it on `PYTHONPATH`. Every value is
invented. The files are not byte-reproducible (zip timestamps), and nothing
compares their bytes: the tests compare what they hold, against
`../uploads/converted_sheet.csv` and against the literals in
`rust/crates/lakehouse-api/src/upload_workbook.rs`.

Written here:

- `stock.xlsx` and `stock.xls`: the same five sheets in each (`Stock`,
  `Quirks`, `Offset`, `Hidden notes` (hidden), `Empty`).
- `formula.xlsx`: formula cells with stored results. openpyxl cannot store a
  result, so the sheet XML is patched after it is written (the script fails
  if the patch does not apply).
- `date1904.xlsx`: dates in the 1904 date system.
- `plain.zip`: a zip that is not a workbook.
"""

from __future__ import annotations

import datetime as dt
import io
import pathlib
import zipfile

import openpyxl
import xlwt
from openpyxl.utils.datetime import CALENDAR_MAC_1904

HERE = pathlib.Path(__file__).parent
FIXED = dt.datetime(2025, 1, 1, 0, 0, 0)

# (label, value, number format). `None` as a value leaves the cell empty.
QUIRKS = [
    ("whole number", 1.0, None),
    ("fraction", 0.1, None),
    ("one third", 1 / 3, None),
    ("large", 1e21, None),
    ("tiny", 1e-7, None),
    ("negative", -42.5, None),
    ("thousands display", 1234.5, "#,##0.00"),
    ("percent", 0.25, "0%"),
    ("date", dt.date(2025, 9, 24), "yyyy-mm-dd"),
    ("date us display", dt.date(2025, 9, 24), "mm/dd/yyyy"),
    ("date time", dt.datetime(2025, 9, 24, 13, 30, 5), "yyyy-mm-dd hh:mm:ss"),
    ("time", dt.time(13, 30, 0), "h:mm:ss"),
    ("duration", 1.5, "[h]:mm:ss"),
    ("bool true", True, None),
    ("bool false", False, None),
    ("error div0", "#DIV/0!", None),
    ("error na", "#N/A", None),
    ("text with comma", "Bolt, hex M6", None),
    ("text with quote", 'She said "ok"', None),
    ("multi-line", "line one\nline two", None),
    ("padded", "  padded  ", None),
    ("empty", None, None),
    ("text number", "007", None),
    ("unicode", "Käse 日本", None),
    ("leading quote", '"quoted" start', None),
]

STOCK = [
    ("sku", "name", "qty", "price", "received"),
    ("A-100", "Bolt, hex M6", 10, 1.5, dt.date(2025, 9, 24)),
    ("A-200", "Washer", 250, 0.05, dt.date(2025, 9, 25)),
    ("A-300", "Flange DN50 – steel", 4, 12, dt.date(2025, 10, 1)),
]


def serial(value: dt.date | dt.datetime | dt.time) -> float:
    """The Excel serial (1900 system) of a date, date-time or time."""
    if isinstance(value, dt.time):
        return (value.hour * 3600 + value.minute * 60 + value.second) / 86400
    if not isinstance(value, dt.datetime):
        value = dt.datetime.combine(value, dt.time())
    delta = value - dt.datetime(1899, 12, 30)
    return delta.days + delta.seconds / 86400


def xlsx_sheets(wb: openpyxl.Workbook) -> None:
    ws = wb.active
    ws.title = "Stock"
    for row in STOCK:
        ws.append(row)
    for cell in ws["E"][1:]:
        cell.number_format = "yyyy-mm-dd"

    ws = wb.create_sheet("Quirks")
    ws.append(("case", "value"))
    for index, (label, value, fmt) in enumerate(QUIRKS, start=2):
        ws.cell(row=index, column=1, value=label)
        if value is not None:
            cell = ws.cell(row=index, column=2, value=value)
            if fmt:
                cell.number_format = fmt
    last = len(QUIRKS) + 2
    ws.cell(row=last, column=1, value="merged top")
    ws.cell(row=last, column=2, value="M")
    ws.cell(row=last + 1, column=1, value="merged bottom")
    ws.merge_cells(start_row=last, end_row=last + 1, start_column=2, end_column=2)

    ws = wb.create_sheet("Offset")
    ws["C3"], ws["D3"] = "id", "label"
    ws["C4"], ws["D4"] = 1, "first"
    ws["C6"], ws["D6"] = 3, "third"

    ws = wb.create_sheet("Hidden notes")
    ws["A1"], ws["A2"] = "note", "internal"
    ws.sheet_state = "hidden"

    wb.create_sheet("Empty")


def make_xlsx() -> None:
    wb = openpyxl.Workbook()
    wb.properties.created = wb.properties.modified = FIXED
    wb.properties.creator = "lakehouse-fixtures"
    xlsx_sheets(wb)
    wb.save(HERE / "stock.xlsx")


def make_xls() -> None:
    wb = xlwt.Workbook()

    def style(fmt: str | None) -> xlwt.XFStyle:
        return xlwt.easyxf(num_format_str=fmt) if fmt else xlwt.XFStyle()

    ws = wb.add_sheet("Stock")
    for r, row in enumerate(STOCK):
        for c, value in enumerate(row):
            if isinstance(value, dt.date):
                ws.write(r, c, serial(value), style("yyyy-mm-dd"))
            else:
                ws.write(r, c, value)

    ws = wb.add_sheet("Quirks")
    ws.write(0, 0, "case")
    ws.write(0, 1, "value")
    for index, (label, value, fmt) in enumerate(QUIRKS, start=1):
        ws.write(index, 0, label)
        if value is None:
            continue
        if isinstance(value, str) and value.startswith("#"):
            # xlwt spells `#N/A` as `#N/A!`.
            ws.row(index).set_cell_error(1, "#N/A!" if value == "#N/A" else value)
        elif isinstance(value, (dt.date, dt.time)):
            ws.write(index, 1, serial(value), style(fmt))
        else:
            ws.write(index, 1, value, style(fmt))
    last = len(QUIRKS) + 1
    ws.write(last, 0, "merged top")
    ws.write_merge(last, last + 1, 1, 1, "M")
    ws.write(last + 1, 0, "merged bottom")

    ws = wb.add_sheet("Offset")
    ws.write(2, 2, "id")
    ws.write(2, 3, "label")
    ws.write(3, 2, 1)
    ws.write(3, 3, "first")
    ws.write(5, 2, 3)
    ws.write(5, 3, "third")

    ws = wb.add_sheet("Hidden notes")
    ws.write(0, 0, "note")
    ws.write(1, 0, "internal")
    ws.visibility = 1

    wb.add_sheet("Empty")
    wb.save(str(HERE / "stock.xls"))


def patch_sheet_xml(src: bytes, edits: list[tuple[str, str]]) -> bytes:
    """Rewrite `xl/worksheets/sheet1.xml` of an `.xlsx` held in memory."""
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(src)) as zin, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                text = data.decode("utf-8")
                for old, new in edits:
                    assert old in text, f"the patch {old!r} did not apply to {text}"
                    text = text.replace(old, new)
                data = text.encode("utf-8")
            zout.writestr(item.filename, data)
    return out.getvalue()


def make_formula() -> None:
    wb = openpyxl.Workbook()
    wb.properties.created = wb.properties.modified = FIXED
    ws = wb.active
    ws.title = "Formula"
    ws["A1"], ws["B1"] = "sum", "=1+2"
    ws["A2"], ws["B2"] = "text", '="a"&"b"'
    ws["A3"], ws["B3"] = "error", "=1/0"
    buffer = io.BytesIO()
    wb.save(buffer)
    raw = buffer.getvalue()
    with zipfile.ZipFile(io.BytesIO(raw)) as z:
        text = z.read("xl/worksheets/sheet1.xml").decode("utf-8")
    empty = "<v></v>" if "<v></v>" in text else "<v />"
    edits = [
        (f"<f>1+2</f>{empty}", "<f>1+2</f><v>3</v>"),
        (f'<f>"a"&amp;"b"</f>{empty}', '<f>"a"&amp;"b"</f><v>ab</v>'),
        (f"<f>1/0</f>{empty}", "<f>1/0</f><v>#DIV/0!</v>"),
    ]
    patched = patch_sheet_xml(raw, edits)
    # Cell types: B2 a string result, B3 an error result.
    with zipfile.ZipFile(io.BytesIO(patched)) as z:
        sheet = z.read("xl/worksheets/sheet1.xml").decode("utf-8")
    import re

    def retype(cell: str, kind: str, xml: str) -> str:
        pattern = re.compile(rf'<c r="{cell}"([^>]*?)(?: t="[a-z]+")?([^>]*)>')
        assert pattern.search(xml), f"no cell {cell}"
        return pattern.sub(lambda m: f'<c r="{cell}"{m.group(1)}{m.group(2)} t="{kind}">', xml, count=1)

    sheet = retype("B2", "str", sheet)
    sheet = retype("B3", "e", sheet)
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(patched)) as zin, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                data = sheet.encode("utf-8")
            zout.writestr(item.filename, data)
    (HERE / "formula.xlsx").write_bytes(out.getvalue())


def make_1904() -> None:
    wb = openpyxl.Workbook()
    wb.epoch = CALENDAR_MAC_1904
    wb.properties.created = wb.properties.modified = FIXED
    ws = wb.active
    ws.title = "Dates"
    ws["A1"] = "when"
    ws["A2"] = dt.date(2025, 9, 24)
    ws["A2"].number_format = "yyyy-mm-dd"
    ws["A3"] = dt.datetime(2025, 9, 24, 13, 30, 5)
    ws["A3"].number_format = "yyyy-mm-dd hh:mm:ss"
    ws["A4"] = dt.time(13, 30, 0)
    ws["A4"].number_format = "h:mm:ss"
    wb.save(HERE / "date1904.xlsx")


def make_zip() -> None:
    with zipfile.ZipFile(HERE / "plain.zip", "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("notes.txt", "not a workbook\n")


if __name__ == "__main__":
    make_xlsx()
    make_xls()
    make_formula()
    make_1904()
    make_zip()
